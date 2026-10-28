"""Small python-docx helpers used by make_results_doc.py and make_judge_doc.py: fonts, inline bold/italic, page numbers."""
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_font(style_or_run, latin, east, size=None, bold=None):
    f = style_or_run.font
    f.name = latin
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(k), latin)
    rfonts.set(qn('w:eastAsia'), east)


def smart(text):
    """Straight quotes to curly quotes."""
    out, open_q = [], True
    for ch in text:
        if ch == '"':
            out.append('“' if open_q else '”'); open_q = not open_q
        elif ch == "'":
            out.append('’')
        else:
            out.append(ch)
    return ''.join(out)


def add_runs(par, text, latin, east):
    """Add text to a paragraph, with **bold** and *italic* marked inline."""
    text = smart(text)
    for tok in re.split(r'(\*\*.+?\*\*|\*.+?\*)', text):
        if not tok:
            continue
        if tok.startswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*'):
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        set_font(r, latin, east)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        run._r.append(el)
