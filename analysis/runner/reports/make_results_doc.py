"""Build the coding results document (Word) from harness.tables output."""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from docx_helpers import set_font, add_runs, add_page_number  # noqa: E402

ap = argparse.ArgumentParser(description=__doc__)
ap.add_argument('--runs', required=True, type=Path, help='run records folder (data_dir in config/paths.local.yaml)')
ap.add_argument('--out', default=Path('.'), type=Path, help='folder for the output')
ARGS = ap.parse_args()

TABLES = ARGS.runs / 'coding/coding_v1-20261024T191027Z/tables/20261026T201743Z/tables.json'
OUT = ARGS.out / 'Pilot_Results_LLM_Coding_2026-10-26.docx'
LATIN = EAST = 'Times New Roman'
BODY = 10.5

t = json.loads(TABLES.read_text(encoding='utf-8'))
L = t['labels']
MODELS = list(t['by_model']['F']['C1'])
SHORT = {'claude-fable-5-1': 'Fable', 'claude-haiku-5-5': 'Haiku', 'gpt-6-astra': 'Astra', 'gpt-6-luna': 'Luna',
         'gemini-4-argon': 'Argon', 'gemini-3-5-flash-lite': 'Flash-Lite', 'deepseek-v4p1-flash': 'DeepSeek',
         'kimi-k3': 'Kimi'}
STATUS = t['five_in_six']
MARK = {c: ('†' if s == 'not reliably coded' else '‡' if s in ('Sonnet only', 'Qwen only') else '') for c, s in STATUS.items()}


def rng(lo, hi, scale=100, dp=0):
    a, b = round(lo * scale, dp), round(hi * scale, dp)
    f = (lambda x: f'{x:.{dp}f}')
    return f(a) if a == b else f'{f(a)}–{f(b)}'


def pc(c):
    return rng(c['lower'], c['higher'])


def per(d, lvl):
    return rng(d[f'{lvl}_lower'], d[f'{lvl}_higher'], 1, 1)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.0)
normal = doc.styles['Normal']; set_font(normal, LATIN, EAST, BODY)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.12
h = doc.styles['Heading 1']; set_font(h, LATIN, EAST, BODY + 1.5, bold=True); h.font.italic = False
h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4); h.paragraph_format.keep_with_next = True
add_page_number(sec)


def P(text, size=None, after=None, keep=False, center=False, style=None):
    p = doc.add_paragraph(style=style); add_runs(p, text, LATIN, EAST)
    for r in p.runs:
        if size:
            r.font.size = Pt(size)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def table(rows, widths, size=8.5, center_from=1, bold_rows=()):
    tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = 'Table Grid'
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), '60'); el.set(qn('w:type'), 'dxa'); mar.append(el)
    tb._tbl.tblPr.append(mar)
    for j, w in enumerate(widths):
        tb.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        tb.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, val in enumerate(row):
            cell = tb.cell(i, j); cell.width = Cm(widths[j])
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            if j >= center_from:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            add_runs(p, str(val), LATIN, EAST)
            for r in p.runs:
                r.font.size = Pt(size)
                if i == 0 or i in bold_rows:
                    r.bold = True
            if i == 0:
                shade(cell, 'E7E7E7')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def caption(text):
    P(text, size=BODY - 0.5, after=3, keep=True)


def note(text):
    P(text, size=BODY - 2, after=8)


ag = {r['concern']: r for r in t['agreement']}
pcF, pcFP = t['planted_clean']['F'], t['planted_clean']['FP']
tp = t['type_prompt']
eff, fmt = t['effort'], t['format']
cons = t['construction']
crit = t['criticisms']
q3 = t['single_coder'].get('C3')

# ---------------- title and reading guide ----------------
def what(text):
    P('*What it shows:* ' + text, size=BODY - 1, after=3, keep=True)


p = doc.add_paragraph(); r = p.add_run('Pilot Results: LLM Coding of the AI Reviews'); set_font(r, LATIN, EAST, BODY + 5, bold=True)
p.paragraph_format.space_after = Pt(2)
P('Investigator 1, 26 October 2026', after=8)
P('Two LLM coders, Claude Sonnet 5.5 and Qwen3.8, coded all 540 AI reviews (480 main, 40 effort, 20 format) with the '
  'frozen coding sheet (v1.0). We both coded the core 24 main reviews by hand to check them. The judge analysis and the '
  'recommendations are in a separate file.')
P('**How to read the tables.** F means the review raised a planted problem and gave the reason in the coding sheet. '
  '“F or P” means it raised the problem at all. When the two LLM coders disagree, we compute each result twice, once '
  'with the lower code and once with the higher one. A cell like 23–28 shows both; a single number means they agree. '
  '† marks C5 and C7, which are not reliably coded (Table 1). ‡ marks C3, which is also shown with Qwen’s codes alone.')

# ---------------- summary ----------------
doc.add_heading('Summary', level=1)
allF, allFP = pcF['All'], pcFP['All']
bmr, anyr = t['by_model_per_review']['F'], t['by_model_any_run']['F']
best = max(MODELS, key=lambda m: bmr[m]['lower'])
worst = min(MODELS, key=lambda m: bmr[m]['lower'])
used, other = cons['F']['used in construction']['flawed'], cons['F']['not used in construction']['flawed']
summary = [
    f'The two LLM coders gave the same code in {ag["All"]["llm_llm"]["agreement"]:.0%} of cells, and the two of us in '
    f'{ag["All"]["human_human"]["agreement"]:.0%}. Against our agreed codes, both LLM coders pass five in six on C1, C2, C4, '
    f'C6 and C8. Only Qwen passes on C3. C5 and C7 are not reliably coded (Table 1).',
    f'A planted problem was raised with its reason in {pc(allF["planted"])}% of cases on the planted version and in '
    f'{pc(allF["clean"])}% on the clean version. Counting any mention, it was {pc(allFP["planted"])}% and {pc(allFP["clean"])}% (Table 4).',
    f'Per planted review, models caught from {rng(bmr[worst]["lower"], bmr[worst]["higher"], 1, 1)} '
    f'({L["models"][worst]}) to {rng(bmr[best]["lower"], bmr[best]["higher"], 1, 1)} ({L["models"][best]}) of the 8 '
    f'problems. Over its 30 runs, each model caught between {min(anyr[m]["lower"] for m in MODELS)} and '
    f'{max(anyr[m]["higher"] for m in MODELS)} at least once (Table 3).',
    f'The prompt made little difference: {pc(tp["F"]["All"]["p1_bare_v1"])}%, {pc(tp["F"]["All"]["p2_reader_v1"])}% and '
    f'{pc(tp["F"]["All"]["p3_criteria_v1"])}% for the three prompts (Table 5).',
    f'Models used in writing the paper caught {rng(used["lower"], used["higher"], 1, 1)} problems per planted review, '
    f'the others {rng(other["lower"], other["higher"], 1, 1)} (Table 6).',
    f'GPT-6 Astra caught about the same at low, default and max effort ({per(eff["low"]["flawed"], "F")}, '
    f'{per(eff["default"]["flawed"], "F")} and {per(eff["max"]["flawed"], "F")}; Table 7). Gemini 4 Argon caught '
    f'{per(fmt["PDF"]["flawed"], "F")} from the PDF and {per(fmt["Text + figures"]["flawed"], "F")} from text with figures (Table 8).',
    f'Reviews made about {crit["all"]["clean"]:.0f} to {crit["all"]["flawed"]:.0f} criticisms each on both versions, from '
    f'{min(v["flawed"] for v in crit["by_model"].values()):.1f} for Gemini 3.5 Flash-Lite to '
    f'{max(v["flawed"] for v in crit["by_model"].values()):.1f} for Claude Haiku 5.5 (Table 9).',
]
for g in summary:
    P('• ' + g, after=3)

# ---------------- 1. coding ----------------
doc.add_heading('1. Can we trust the coding?', level=1)
caption('**Table 1.** How well do the coders agree? (first codes)')
what('Agreement between the two LLM coders, between the two of us, and between each of us and each LLM coder. The last '
     'column applies the five-in-six rule: an LLM coder must match our agreed codes in at least 20 of the 24 reviews.')
rows = [['Concern', 'LLM coders (540): same / κ / AC1', 'Inv. 1 and Inv. 2 (24): same / κ / AC1', 'Inv. 1–Sonnet', 'Inv. 1–Qwen',
         'Inv. 2–Sonnet', 'Inv. 2–Qwen', 'Sonnet–agreed', 'Qwen–agreed', 'Five in six']]
f2 = lambda x: '–' if x is None else f'{x:.2f}'
for c in list(L['titles']) + ['All']:
    a = ag[c]
    rows.append([c + MARK.get(c, ''), f'{a["llm_llm"]["agreement"]:.0%} / {f2(a["llm_llm"]["kappa"])} / {f2(a["llm_llm"]["ac1"])}',
                 f'{a["human_human"]["agreement"]:.0%} / {f2(a["human_human"]["kappa"])} / {f2(a["human_human"]["ac1"])}',
                 a['coder1_sonnet']['same'], a['coder1_qwen']['same'], a['coder2_sonnet']['same'], a['coder2_qwen']['same'],
                 a['sonnet_agreed']['same'], a['qwen_agreed']['same'],
                 {'both coders pass': 'Both pass', 'Qwen only': 'Qwen only', 'Sonnet only': 'Sonnet only',
                  'not reliably coded': 'Neither'}.get(STATUS.get(c), '')])
table(rows, [1.4, 2.6, 2.6, 1.15, 1.15, 1.15, 1.15, 1.3, 1.3, 1.6], size=8, bold_rows=(len(rows) - 1,))
note('Same: share of identical codes. κ: Cohen’s kappa. AC1: Gwet’s AC1. The middle columns count identical codes out of '
     '24 reviews (192 in the All row). We disagreed in 55 cells, and Inv. 1 took Inv. 2’s code in all of them, so the agreed codes '
     'are Inv. 2’s codes.')

pv = t['provider_check']
caption('**Table 2.** Does Sonnet go easy on reviews by Claude models? (share of cells coded, main reviews)')
what('Sonnet comes from the same company as two of the reviewers (Claude Fable 5.1 and Haiku 5.5). If it favored them, it '
     'would code their reviews higher than Qwen does, by more than it does for the other reviews.')
rows = [['Reviews', 'Reviews (n)', 'Sonnet: F', 'Qwen: F', 'Sonnet: F or P', 'Qwen: F or P']]
for g, v in pv.items():
    rows.append([g, v['F']['n_reviews'], f'{v["F"]["sonnet"]:.0%}', f'{v["F"]["qwen"]:.0%}', f'{v["FP"]["sonnet"]:.0%}', f'{v["FP"]["qwen"]:.0%}'])
table(rows, [3.6, 1.8, 2.0, 2.0, 2.4, 2.4])

# ---------------- 2. by model ----------------
doc.add_heading('2. Which problems did each model catch?', level=1)
caption('**Table 3.** Planted version: % of each model’s 30 runs that raised the problem with its reason (F)')
what('Which problems each model finds. The last two rows show how many problems a model catches in one run, and how many '
     'it catches at least once over its 30 runs.')
rows = [['Concern', 'Type'] + [SHORT[m] for m in MODELS]]
for c in L['titles']:
    rows.append([c + MARK[c], L['types'][c]] + [pc(t['by_model']['F'][c][m]) for m in MODELS])
if q3:
    rows.append(['C3‡ (Qwen)', L['types']['C3']] + [f'{q3["by_model"]["F"][m]["share"]:.0%}'.rstrip('%') for m in MODELS])
rows.append(['Per run (of 8)', ''] + [rng(bmr[m]['lower'], bmr[m]['higher'], 1, 1) for m in MODELS])
rows.append(['At least once in 30 runs (of 8)', ''] + [rng(anyr[m]['lower'], anyr[m]['higher'], 1) for m in MODELS])
table(rows, [2.6, 2.2] + [1.42] * 8, size=8, center_from=2, bold_rows=(len(rows) - 2, len(rows) - 1))
note('30 runs per model: 3 prompts × 10. Fable, Haiku = Claude Fable 5.1, Haiku 5.5; Astra, Luna = GPT-6 Astra, Luna; '
     'Argon = Gemini 4 Argon; Flash-Lite = Gemini 3.5 Flash-Lite; DeepSeek = DeepSeek V4.1 Flash; Kimi = Kimi K3.')

# ---------------- 3. planted vs clean ----------------
doc.add_heading('3. Do the models raise a problem because it is there?', level=1)
caption('**Table 4.** % of main reviews that raised each problem, planted vs clean version (240 reviews each)')
what('In the clean version the planted problem is fixed. If a model raises a problem about as often there, it says it '
     'whether or not the problem is there.')
rows = [['Concern', 'Type', 'F: planted', 'F: clean', 'F or P: planted', 'F or P: clean']]
for c in list(L['titles']) + ['All']:
    rows.append([c + MARK.get(c, ''), L['types'].get(c, 'all 8'), pc(pcF[c]['planted']), pc(pcF[c]['clean']),
                 pc(pcFP[c]['planted']), pc(pcFP[c]['clean'])])
if q3:
    rows.insert(4, ['C3‡ (Qwen)', L['types']['C3']] + [f'{q3["planted_clean"][l][v]["share"]:.0%}'.rstrip('%')
                                                         for l in ('F', 'FP') for v in ('planted', 'clean')])
table(rows, [2.3, 2.6, 2.4, 2.4, 2.6, 2.6], center_from=2, bold_rows=(len(rows) - 1,))
note('A criticism on the clean version is “raised on clean”. It may still be a fair point for that version (coding sheet, '
     'section 1, rule 4).')

# ---------------- 4. type x prompt ----------------
doc.add_heading('4. Does the prompt matter?', level=1)
caption('**Table 5.** Planted version: % of problems raised, by problem type and prompt')
what('Whether a fuller prompt helps the models find more, and whether it helps more for some kinds of problems.')
rows = [['Type', 'Concerns', 'F: Bare', 'F: Reader', 'F: Reader + criteria', 'F or P: Bare', 'F or P: Reader', 'F or P: Reader + criteria']]
groups = {ty: [c for c in L['titles'] if L['types'][c] == ty] for ty in ('Technical', 'Technical/mixed', 'Mixed', 'Taste/mixed')}
for ty in list(groups) + ['All']:
    cs = ', '.join(c + MARK[c] for c in groups[ty]) if ty != 'All' else 'all 8'
    rows.append([ty, cs] + [pc(tp[l][ty][p]) for l in ('F', 'FP') for p in L['prompts']])
table(rows, [2.3, 2.3, 1.7, 1.7, 2.0, 1.8, 1.9, 2.3], size=8, center_from=2, bold_rows=(len(rows) - 1,))
note('80 planted-version reviews per prompt (8 models × 10). Types are the labels from the construction register.')

# ---------------- 5. construction ----------------
doc.add_heading('5. Do the models used in writing the paper do better?', level=1)
caption('**Table 6.** Problems raised per review (of 8): models used in writing the paper vs the others')
what('Claude and GPT models helped build the manuscript. This shows whether they catch more than the other models.')
rows = [['Models', 'F: planted', 'F: clean', 'F or P: planted', 'F or P: clean']]
for tag, lab in (('used in construction', 'Used in writing (Claude, GPT-6)'), ('not used in construction', 'Not used (Gemini, DeepSeek, Kimi)')):
    rows.append([lab] + [rng(cons[l][tag][v]['lower'], cons[l][tag][v]['higher'], 1, 1) for l in ('F', 'FP') for v in ('flawed', 'clean')])
table(rows, [5.4, 2.4, 2.4, 2.6, 2.6])
note('4 models in each group, 30 reviews per model and version.')

# ---------------- 6. effort and format ----------------
doc.add_heading('6. Do effort and input format matter?', level=1)
caption('**Table 7.** GPT-6 Astra on prompt 3: problems raised per review (of 8) and criticisms, by effort')
what('Whether more reasoning finds more problems, or only adds more criticism.')
rows = [['Effort', 'F: planted', 'F: clean', 'F or P: planted', 'F or P: clean', 'Criticisms: planted', 'Criticisms: clean']]
for k, lab in (('low', 'Low'), ('default', 'Default (main runs)'), ('max', 'Max')):
    d = eff[k]
    rows.append([lab, per(d['flawed'], 'F'), per(d['clean'], 'F'), per(d['flawed'], 'FP'), per(d['clean'], 'FP'),
                 f'{d["flawed"]["criticisms"]:.1f}', f'{d["clean"]["criticisms"]:.1f}'])
table(rows, [3.2, 1.9, 1.9, 2.2, 2.2, 2.3, 2.3])
caption('**Table 8.** Gemini 4 Argon on prompt 2: problems raised per review (of 8) and criticisms, by input format')
what('Whether reading the PDF or the text with figures changes what the model finds.')
rows = [['Input', 'F: planted', 'F: clean', 'F or P: planted', 'F or P: clean', 'Criticisms: planted', 'Criticisms: clean']]
for k in ('PDF', 'Text + figures'):
    d = fmt[k]
    rows.append([k, per(d['flawed'], 'F'), per(d['clean'], 'F'), per(d['flawed'], 'FP'), per(d['clean'], 'FP'),
                 f'{d["flawed"]["criticisms"]:.1f}', f'{d["clean"]["criticisms"]:.1f}'])
table(rows, [3.2, 1.9, 1.9, 2.2, 2.2, 2.3, 2.3])
note('10 reviews per version in each row.')

# ---------------- 7. criticisms ----------------
doc.add_heading('7. How much do the reviews criticize?', level=1)
caption('**Table 9.** Criticisms per review (average of the two LLM coders’ counts)')
what('How much each model criticizes, right or wrong. More criticism is not the same as catching more problems.')
rows = [['Model or prompt', 'Planted', 'Clean']]
for m in MODELS:
    rows.append([L['models'][m], f'{crit["by_model"][m]["flawed"]:.1f}', f'{crit["by_model"][m]["clean"]:.1f}'])
for p, lab in L['prompts'].items():
    rows.append([f'Prompt: {lab}', f'{crit["by_prompt"][p]["flawed"]:.1f}', f'{crit["by_prompt"][p]["clean"]:.1f}'])
rows.append(['All main reviews', f'{crit["all"]["flawed"]:.1f}', f'{crit["all"]["clean"]:.1f}'])
table(rows, [6.0, 2.6, 2.6], bold_rows=(len(rows) - 1,))

doc.save(OUT)
print(OUT)
