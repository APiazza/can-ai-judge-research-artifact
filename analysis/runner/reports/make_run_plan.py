"""Build the AI review run plan (Word). Run make_workflow_fig.py first, into the same folder."""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from docx_helpers import set_font, add_runs, add_page_number  # noqa: E402

ap = argparse.ArgumentParser(description=__doc__)
ap.add_argument('--out', default=Path('.'), type=Path, help='folder with the figure and for the Word file')
ARGS = ap.parse_args()
LATIN = 'Times New Roman'

CRITERIA = ['Importance of Problem', 'Relevance to Organization Science', 'Incorporation of Relevant Literature',
            'Development of Theory', 'Appropriateness of Design', 'Appropriateness of Analysis and Interpretation',
            'Legitimacy of Conclusions', 'Writing Clarity and Organization']

EN = {
    'out': 'AI_Run_Plan_2026-10-26.docx', 'east': 'Times New Roman', 'body': 10.5,
    'fig': 'workflow_2026-10-24.png', 'figw': 13.6,
    'title': 'AI Review Runs: Plan',
    'byline': 'Investigator 1, 26 October 2026',
    'intro': 'This plan sets out the materials, models, prompts, runs, timeline and coding for the AI review runs of the pilot, '
             'and the exploratory judge analysis.',
    'h1': '1. Workflow',
    'workflow': 'Figure 1 shows the steps. Runs start only after the plan is approved, and the coding sheet is frozen before '
                'any main review is read. The answer key is used only for coding, and reader reports enter only after IRB approval. '
                'The judge analysis (section 8) was added after the coding sheet was frozen and is exploratory.',
    'figcap': '**Figure 1.** Workflow. Dashed boxes are kept apart from the runs.',
    'h2': '2. Materials',
    'materials': [
        'The AI copies of pilot-v6 and clean-matched-v6 are in corpus/papers/STATUS-P01/ai-run-inputs/, in the subfolders '
        'pilot-v6/ and clean-matched-v6/. The page-1 notice is removed and every other page is unchanged. Both files are named '
        'manuscript.pdf and have no document properties. The abstracts differ only in the two sentences that carry C3 and C6. '
        'The run code reads only this folder.',
        'Models that accept PDFs receive manuscript.pdf, which keeps each figure next to its text. Models without PDF input '
        'receive manuscript.md with each figure placed at its position in the text, as in the PDF. In every request the '
        'manuscript comes first and the prompt after it.',
        'The answer key is in corpus/keys/STATUS-P01-pilot-v6-key.json, with a readable .md version. It refers to pilot-v6 and '
        'gives the matching passage in clean-matched-v6 for each concern. It is outside what the run code can read and is never '
        'sent to a model under review. The build record is in corpus/prompt_history/STATUS-P01-AI-COPIES-001/.',
    ],
    'h3': '3. Models',
    'models': [
        ['Provider', 'Model', 'API ID', 'Served by', 'Input'],
        ['Anthropic', 'Claude Fable 5.1 (flagship)', 'claude-fable-5-1', 'Anthropic API', 'PDF'],
        ['Anthropic', 'Claude Haiku 5.5 (small)', 'claude-haiku-5-5', 'Anthropic API', 'PDF'],
        ['OpenAI', 'GPT-6 Astra (flagship)', 'gpt-6-astra', 'OpenAI API', 'PDF'],
        ['OpenAI', 'GPT-6 Luna (small)', 'gpt-6-luna', 'OpenAI API', 'PDF'],
        ['Google', 'Gemini 4 Argon (flagship)', 'gemini-4-argon', 'Gemini API', 'PDF'],
        ['Google', 'Gemini 3.5 Flash-Lite (small)', 'gemini-3.5-flash-lite', 'Gemini API', 'PDF'],
        ['Open-weight', 'DeepSeek V4.1 Flash', 'accounts/fireworks/models/deepseek-v4p1-flash', 'Fireworks', 'Text + figures'],
        ['Open-weight', 'Kimi K3', 'accounts/fireworks/models/kimi-k3', 'Fireworks', 'Text + figures'],
    ],
    'model_widths': [2.3, 3.7, 5.0, 2.6, 2.6],
    'models_text': [
        'PDF input for Gemini 4 Argon is confirmed in the first practice run.',
        'DeepSeek V4 Pro is available on Fireworks only as a dedicated deployment, so DeepSeek V4.1 Flash, DeepSeek’s newest '
        'model (released 10 September 2026), is used instead. DeepSeek V4.1 Flash and Kimi K3 run on Fireworks’ serverless API '
        'and read the text version with each figure at its place. Kimi K3 is served with MXFP4 weights and MXFP8 activations. '
        'Fireworks does not state the serving precision of DeepSeek V4.1 Flash; its released weights are BF16 and FP8.',
        'In the main runs, models whose provider does not report effort receive their documented default explicitly: high '
        'for Claude Fable 5.1, medium for Claude Haiku 5.5 and Gemini 4 Argon, and minimal for Gemini 3.5 Flash-Lite. GPT-6 '
        'Astra and Luna (medium, reported by OpenAI) and DeepSeek V4.1 Flash (high, reported by Fireworks) are left at their '
        'defaults, as is Kimi K3, whose default is not documented. A check of 10 practice calls with these values found no '
        'material change in token counts.',
        'The manuscript was built with help from Claude and Codex, so the Anthropic and OpenAI models are tagged as used in '
        'construction and the Google and open-weight models as not used in construction. The tag is stored with every run, '
        'so the results can be split by it.',
    ],
    'h4': '4. Prompts',
    'p1': '**Prompt 1 (Bare):** “Review this paper.”',
    'p2_label': '**Prompt 2 (Reader instructions)**',
    'p2': 'Please read the attached paper and write the referee report and recommendation you would ordinarily provide to a '
          'journal.\n\nAt the end of your report, state your recommendation to the editor: Reject, Major revision, Minor '
          'revision, or Accept.',
    'p3_label': '**Prompt 3 (Reader instructions plus criteria)**',
    'p3': 'Please read the attached paper and write the referee report and recommendation you would ordinarily provide to a '
          'journal.\n\nIn preparing your report, please consider the following criteria. A manuscript need not have '
          'weaknesses on every criterion, and your report need not address each one.\n\n'
          + '\n'.join(f'{i}. {c}' for i, c in enumerate(CRITERIA, 1)) +
          '\n\nAt the end of your report, state your recommendation to the editor: Reject, Major revision, Minor revision, '
          'or Accept.',
    'prompts_text': 'Readers received no separate review instructions; their task appears only in the study information on '
                    'the consent page. Prompt 2 uses the sentence that states the task, with the recommendation options from '
                    'survey item Q2.1. Prompt 3 adds the eight criteria of the Organization Science reviewer scorecard, without '
                    'ratings. Readers knew that the paper was simulated, that problems had been introduced and that this was a '
                    'research task; the models receive an ordinary request for a referee report. If anything, the comparison '
                    'favors the readers.',
    'h5': '5. Runs',
    'runs': [
        ['Set', 'Manuscript', 'Prompts', 'Runs'],
        ['Main', 'Pilot-v6', 'Prompts 1–3', '8 models × 3 prompts × 10 = 240'],
        ['Main', 'Clean-matched-v6', 'Prompts 1–3', '8 models × 3 prompts × 10 = 240'],
        ['Effort', 'Both', 'Prompt 3', 'GPT-6 Astra at low and max: 2 × 2 × 10 = 40'],
        ['Format', 'Both', 'Prompt 2', 'Gemini 4 Argon on text + figures: 2 × 10 = 20'],
        ['Practice', 'Both', 'Prompts 1–3', '20'],
        ['Total', '', '', '560'],
    ],
    'runs_text': [
        'Each run is a new API call with no tools, no web access and default temperature. Each run record holds the model ID, '
        'construction tag, date, effort, input format, token counts and stop reason. Anthropic requests use prompt caching; '
        'OpenAI and Gemini cache repeated input automatically.',
        'Running the effort comparison on both versions shows whether higher effort finds more problems or only produces more '
        'criticism. The format runs compare Gemini 4 Argon on the PDF and on the text-plus-figures input, to show whether the '
        'input format matters. The practice runs are used only to build the coding sheet and are not reported.',
        'After the first practice runs, the input and output tokens per run and the spend so far by account are reported, '
        'with a cost estimate before the main runs start.',
    ],
    'h6': '6. Timeline',
    'cost': 'Rows up to 25 October show when each step was done; the rows from 26 October are planned.',
    'schedule': [
        ['Date', 'Step'],
        ['15–16 Oct', 'Practice runs and checks'],
        ['18 Oct', 'Main, effort and format runs'],
        ['19–23 Oct', 'Coding sheet drafted on the practice reviews, reviewed, revised and frozen as v1.0 (23 Oct)'],
        ['24–25 Oct', 'LLM coding; human check of the core 24 coded by both investigators'],
        ['26 Oct', 'Agreed codes for the human-check disagreements; judge runs (exploratory)'],
        ['26–27 Oct', 'Results tables; results for the proposal by 27 Oct'],
    ],
    'h7': '7. Coding',
    'coding': [
        'The coding sheet gives F, P or – for each concern, with examples of F, P and a near miss, a rule for a point that '
        'touches two concerns, and a rule for counting criticisms. It was drafted on the practice reviews, reviewed and '
        'revised, and frozen as v1.0 on 23 October, before any main review was read.',
        'Two LLM coders code every review: Claude Sonnet 5.5 and Qwen3.8-2.4T-A95B, the largest open-weight Qwen model '
        '(accounts/fireworks/models/qwen3p8-2p4t-a95b on Fireworks; Fireworks does not state its serving precision, so the '
        'provider’s model information is recorded with each batch). Neither is among the evaluated models. Agreement between '
        'the two coders is reported. Both coders see the key, so only APIs that do not train on submitted data are used. '
        'Model, prompt and version labels are removed and the order is shuffled; blinding is only partial, since quotations '
        'can reveal the version.',
        'The two investigators code the same main reviews: 40 if time allows, and at least a core of 24. The core 24 are three '
        'from each model, one for each prompt; the other 16 are two from each model, one for each version. All 40 were drawn '
        'at random at blinding, before anyone read them, and the core 24 are coded first. Each investigator codes alone, '
        'without seeing the other’s codes or the LLM codes. Agreement is reported between the two investigators and between '
        'each investigator and each LLM coder, with a check that the coders do not treat their own provider’s reviews '
        'differently.',
        'Where the two LLM coders give the same code, that is the final code. Where they differ, no one settles it by hand: '
        'each result is reported with the lower and with the higher of the two codes, and a conclusion is stated only if it '
        'holds under both. The number of criticisms is the average of the two LLM coders’ counts. On the shared reviews, the '
        'two investigators agree on a code for each disagreement, and the principal investigator decides if they cannot agree. These agreed '
        'codes check how accurate each LLM coder is and replace no LLM code. If an LLM coder agrees with them on a concern in '
        'fewer than five in six shared reviews (20 of 24, or 34 of 40), results for that concern are also reported with the '
        'other LLM coder alone; if both fall below this, the concern is reported as not reliably coded.',
        'Under the revised consent, readers can decline AI processing of their reports. Once the IRB approves, the LLM coders '
        'code only the reports of readers who permit it; the rest are coded by hand.',
    ],
    'h8': '8. Judge analysis (exploratory)',
    'judge': [
        'This analysis was added on 24 October, after the coding sheet was frozen, following King and Worthington (2026). It '
        'asks whether a judge’s preference between two reviews tracks which review catches the planted problems, and it is '
        'reported as exploratory.',
        '**Pairs.** Within each version × prompt cell, the 80 main reviews are paired at random so that the two reviews of a '
        'pair come from different models: 40 pairs per cell, 240 in all. The reviews keep their blind IDs from the coding, and '
        'the seed was recorded before pairing. Reader pairs are left out of the pilot.',
        '**Prompt.** King and Worthington’s overall pairwise ranking, reviews only (supplement, Appendix D(b)), word for word, '
        'with their abbreviated reviewer guidelines (Appendix B) inserted where they mark it, is sent as the system message. '
        'The two reviews follow in <review_a> and <review_b> tags. The judge sees no manuscript and no model, version or '
        'prompt label.',
        '**Judges.** GPT-5.5 at maximum effort, as in their study, and Qwen3.8, from a provider whose models were not reviewed, '
        'at the provider default. Each pair is judged in both orders, and a win counts only if the same review wins both times. '
        'This comes to 960 calls (240 pairs × 2 orders × 2 judges), about $120 for GPT-5.5 and $50 for Qwen.',
        '**Analysis.** The verdicts are compared with the LLM codes, with F only and with F or P, and with the lower and the '
        'higher code where the two coders differ: (1) among planted-version pairs whose numbers of planted problems caught '
        'differ, how often the judge prefers the review that catches more; (2) for each planted problem, among pairs where '
        'exactly one review catches it, how often the judge prefers the review that misses it; (3) among clean-version pairs '
        'whose numbers of planted concerns raised differ, how often the judge prefers the review that raises more. Outcome '
        '(1) is read with bands set before any results: 80 percent or more means preference largely tracks detection, '
        '60 percent or less means it is a weak guide, and in between is mixed. Because GPT-5.5 shares a provider with GPT-6 '
        'Astra and Luna, every result is also reported without pairs that contain an OpenAI review. Because Qwen3.8 is also '
        'one of the two coders, the Qwen judge’s results are also reported against Sonnet’s codes alone. The results are '
        'read after the human check.',
        '**Recommendations.** Prompts 2 and 3 asked each reviewer to state Reject, Major revision, Minor revision or Accept. '
        'The stated recommendation is extracted from each of these reviews automatically. Where a review names more than one '
        'option, the first option in its final recommendation is taken, and “accept with minor/major revisions” counts as '
        'minor/major revision. The distributions are compared between the planted and clean versions, by model and prompt. '
        'Prompt 1 reviews are included only if they state a recommendation unprompted, and are reported separately. '
        'This needs no new calls.',
    ],
    'h9': '9. Data',
    'data': 'Only code and prompts go on GitHub. The manuscripts, the key and the model outputs stay in the Dropbox workspace. No '
            'reader data go to any model provider until the IRB approves, and then only the reports of readers who permit it '
            '(section 7). API keys are kept out of the repository and used only for this project’s runs and coding.',
}



def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def build(S):
    east, body = S['east'], S['body']
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.4); sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles['Normal']; set_font(normal, LATIN, east, body)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
    h = doc.styles['Heading 1']; set_font(h, LATIN, east, body + 1.5, bold=True); h.font.italic = False
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(4); h.paragraph_format.keep_with_next = True
    add_page_number(sec)

    def P(text, size=None, keep=False, after=None, center=False):
        p = doc.add_paragraph(); add_runs(p, text, LATIN, east)
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
        if keep:
            p.paragraph_format.keep_with_next = True
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def table(rows, widths, center_from=99, bold_last=False):
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        mar = OxmlElement('w:tblCellMar')
        for side in ('left', 'right'):
            el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), '80'); el.set(qn('w:type'), 'dxa'); mar.append(el)
        t._tbl.tblPr.append(mar)
        for j, w in enumerate(widths):
            t.columns[j].width = Cm(w)
        for i, row in enumerate(rows):
            t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j, val in enumerate(row):
                cell = t.cell(i, j); cell.width = Cm(widths[j])
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                if j >= center_from:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if i < len(rows) - 1:
                    p.paragraph_format.keep_with_next = True
                add_runs(p, val, LATIN, east)
                for r in p.runs:
                    r.font.size = Pt(body - 1)
                    if i == 0 or (bold_last and i == len(rows) - 1):
                        r.bold = True
                if i == 0:
                    shade(cell, 'E7E7E7')
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def prompt_box(text):
        t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cell = t.cell(0, 0); cell.width = Cm(16.2); shade(cell, 'F4F4F4')
        p = cell.paragraphs[0]
        for k, para in enumerate(text.split('\n\n')):
            if k:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.0
            lines = para.split('\n')
            for m, line in enumerate(lines):
                r = p.add_run(line); set_font(r, LATIN, east, body - 1)
                if m < len(lines) - 1:
                    r.add_break()
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph(); r = p.add_run(S['title']); set_font(r, LATIN, east, body + 5, bold=True)
    p.paragraph_format.space_after = Pt(2)
    P(S['byline'], after=10)
    P(S['intro'])

    doc.add_heading(S['h1'], level=1)
    P(S['workflow'])
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True; pic.paragraph_format.space_after = Pt(2)
    pic.add_run().add_picture(str(ARGS.out / S['fig']), width=Cm(S['figw']))
    P(S['figcap'], size=body - 1.5, center=True)

    doc.add_heading(S['h2'], level=1)
    for t in S['materials']:
        P(t)

    doc.add_heading(S['h3'], level=1)
    table(S['models'], S.get('model_widths', [2.4, 6.4, 3.0, 4.4]), center_from=S.get('model_center', 99))
    for t in S['models_text']:
        P(t)

    doc.add_heading(S['h4'], level=1)
    P(S['p1'])
    P(S['p2_label'], keep=True, after=2)
    prompt_box(S['p2'])
    P(S['p3_label'], keep=True, after=2)
    prompt_box(S['p3'])
    P(S['prompts_text'])

    doc.add_heading(S['h5'], level=1)
    table(S['runs'], [2.2, 3.6, 3.4, 7.0], bold_last=True)
    for t in S['runs_text']:
        P(t)

    if S.get('cost'):
        doc.add_heading(S['h6'], level=1)
        P(S['cost'], keep=True)
        table(S['schedule'], [2.6, 13.6])

    doc.add_heading(S['h7'], level=1)
    for t in S['coding']:
        P(t)

    doc.add_heading(S['h8'], level=1)
    for t in S['judge']:
        P(t)

    doc.add_heading(S['h9'], level=1)
    P(S['data'])

    out = ARGS.out / S['out']
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build(EN)
