"""Export the human-check codes to one JSON file for harness.tables (stdlib plus openpyxl; no harness imports).

Reads both coders' completed forms and the disagreement sheet with the agreed codes. Writes, once:
{"sources": {name: {file, sha256}}, "coder1": [...], "coder2": [...], "agreed": [...]} with one entry per review and concern.
First codes are taken as entered (for an unsure code, the provisional one).

    python tools/export_human_codes.py CODER1.xlsx CODER2.xlsx DISAGREEMENTS.xlsx OUT.json
"""
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

CODES = {"F", "P", "–"}


def form(path: Path) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True)["Codes"]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1]:
            if r[3] not in CODES:
                raise SystemExit(f"{path.name}: {r[1]} {r[2][:2]} has no valid code")
            out[(r[1], r[2][:2])] = r[3]
    return out


def main():
    c1p, c2p, dp, outp = map(Path, sys.argv[1:5])
    if outp.exists():
        raise SystemExit(f"{outp} exists; records are written once.")
    c1, c2 = form(c1p), form(c2p)
    if set(c1) != set(c2):
        raise SystemExit("The two forms cover different reviews or concerns.")
    ws = openpyxl.load_workbook(dp, data_only=True)["Disagreements"]
    settled = {(r[1], r[2][:2]): r[10] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}
    agreed = {}
    for k in c1:
        agreed[k] = c1[k] if c1[k] == c2[k] else settled.get(k)
        if agreed[k] not in CODES:
            raise SystemExit(f"No agreed code for {k}")
    if set(settled) != {k for k in c1 if c1[k] != c2[k]}:
        raise SystemExit("The disagreement sheet does not list exactly the cells where the coders differ.")
    rows = lambda d: [{"blind_id": b, "concern": c, "code": d[(b, c)]} for b, c in sorted(d)]
    data = {"sources": {n: {"file": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                        for n, p in (("coder1", c1p), ("coder2", c2p), ("agreed", dp))},
            "coder1": rows(c1), "coder2": rows(c2), "agreed": rows(agreed)}
    outp.parent.mkdir(parents=True, exist_ok=True)
    outp.write_bytes(json.dumps(data, indent=1, ensure_ascii=False).encode("utf-8"))
    print(f"{outp}: {len(c1)} cells per coder, {len(settled)} settled disagreements")


if __name__ == "__main__":
    main()
